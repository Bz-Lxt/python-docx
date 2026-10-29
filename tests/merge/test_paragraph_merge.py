# pyright: reportPrivateUsage=false

"""Integration tests for three-way merging of body paragraphs."""

from __future__ import annotations

from docx.merge import MergeConflict, MergeResult, merge_documents

from .helpers import (
    branches,
    delete_paragraph,
    fork,
    paragraph_texts,
    text_document,
)


class DescribeParagraphMerge:
    def it_combines_nonoverlapping_edit_delete_and_append(self):
        base = text_document("one", "two", "three", "four")
        left, right = branches(base)

        left.paragraphs[0].text = "ONE-left"
        delete_paragraph(left, 2)
        left.add_paragraph("left-tail")
        right.paragraphs[1].text = "TWO-right"
        right.add_paragraph("right-tail")

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        assert paragraph_texts(result.document) == [
            "ONE-left",
            "TWO-right",
            "four",
            "left-tail",
            "right-tail",
        ]

    def it_merges_edits_to_adjacent_paragraphs_without_a_shared_anchor(self):
        base = text_document("a", "b", "c")
        left, right = branches(base)

        left.paragraphs[0].text = "A"
        left.paragraphs[2].text = "C"
        right.paragraphs[1].text = "B"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        assert paragraph_texts(result.document) == ["A", "B", "C"]

    def it_keeps_an_identical_appendix_on_both_branches_once(self):
        base = text_document("keep")
        left, right = branches(base)

        left.add_paragraph("same-insert")
        right.add_paragraph("same-insert")

        result = merge_documents(base, fork(left), fork(right))

        assert paragraph_texts(result.document) == ["keep", "same-insert"]

    def it_keeps_different_insertions_at_the_same_point_from_both_branches(self):
        base = text_document("anchor")
        left, right = branches(base)

        left_new = left.add_paragraph("before?")  # -- move to before "anchor" --
        left.paragraphs[0]._p.addprevious(left_new._p)
        right_new = right.add_paragraph("different")
        right.paragraphs[0]._p.addprevious(right_new._p)

        result = merge_documents(base, fork(left), fork(right))

        assert paragraph_texts(result.document) == ["before?", "different", "anchor"]

    def it_reports_an_overlapping_edit_with_location_branches_and_text(self):
        base = text_document("hello")
        left, right = branches(base)

        left.paragraphs[0].text = "hola"
        right.paragraphs[0].text = "bonjour"

        result = merge_documents(
            base, fork(left), fork(right), left_name="alice", right_name="bob"
        )

        assert isinstance(result, MergeResult)
        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert isinstance(conflict, MergeConflict)
        assert conflict.location == "body[0]"
        assert conflict.kind == "paragraph"
        assert conflict.branches == ("alice", "bob")
        assert conflict.base_text == "hello"
        assert conflict.left_text == "hola"
        assert conflict.right_text == "bonjour"

    def it_writes_both_alternatives_into_conflict_marker_paragraphs(self):
        base = text_document("hello")
        left, right = branches(base)

        left.paragraphs[0].text = "hola"
        right.paragraphs[0].text = "bonjour"

        result = merge_documents(
            base, fork(left), fork(right), left_name="alice", right_name="bob"
        )

        assert paragraph_texts(result.document) == [
            "<<<<<<< alice",
            "hola",
            "=======",
            "bonjour",
            ">>>>>>> bob",
        ]

    def it_reports_delete_against_edit_as_a_conflict_with_empty_side_text(self):
        base = text_document("doomed")
        left, right = branches(base)

        delete_paragraph(left, 0)
        right.paragraphs[0].text = "saved"

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert conflict.kind == "paragraph"
        assert conflict.left_text == ""
        assert conflict.right_text == "saved"

    def it_uses_base_indices_in_conflict_locations(self):
        base = text_document("zero", "one", "two")
        left, right = branches(base)

        left.paragraphs[1].text = "LEFT-ONE"
        right.paragraphs[1].text = "RIGHT-ONE"

        result = merge_documents(base, fork(left), fork(right))

        assert result.conflicts[0].location == "body[1]"

    def it_does_not_mutate_the_input_documents(self):
        base = text_document("x")
        left, right = branches(base)
        left.paragraphs[0].text = "L"
        right.paragraphs[0].text = "R"
        before = {name: doc._element.xml for name, doc in (
            ("base", base), ("left", left), ("right", right)
        )}

        merge_documents(base, left, right)

        assert base._element.xml == before["base"]
        assert left._element.xml == before["left"]
        assert right._element.xml == before["right"]
