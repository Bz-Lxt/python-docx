# pyright: reportPrivateUsage=false

"""Integration tests for cross-package reference remapping during a three-way merge.

Pictures, hyperlinks, numbering definitions and styles all carry package-local
identifiers (rId, numId, styleId). Two branches routinely reuse the *same* source
identifier for *different* content; the merge must import every reference into the
result package so that pictures never "cross the wires" and links/numbering/styles
keep working after a save and reopen.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.merge import merge_documents
from docx.oxml.ns import qn
from docx.shared import Pt

from .helpers import (
    add_hyperlink_paragraph,
    add_numbered_paragraph,
    branches,
    fork,
    image_path,
    numbering_format_for,
    paragraph_texts,
    zip_part_bytes,
)


def empty_base() -> Document:
    return fork(Document())


def image_blob(name: str) -> bytes:
    return Path(image_path(name)).read_bytes()


class DescribeImageReferenceMerge:
    def it_keeps_two_different_pictures_with_the_same_source_rId_distinct(self):
        base = empty_base()
        left, right = branches(base)

        left.add_picture(image_path("python-powered.png"))
        right.add_picture(image_path("monty-truth.png"))
        # -- both branches allocated the very same rId in their own packages --
        left_rid = left._element.body.xpath(".//a:blip/@r:embed")[0]
        right_rid = right._element.body.xpath(".//a:blip/@r:embed")[0]
        assert left_rid == right_rid

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        doc = fork(result.document)
        rids = doc._element.body.xpath(".//a:blip/@r:embed")
        assert len(rids) == 2
        assert len(set(rids)) == 2
        blobs = [doc.part.rels[r_id].target_part.blob for r_id in rids]
        assert blobs == [
            image_blob("python-powered.png"),
            image_blob("monty-truth.png"),
        ]

    def it_shares_one_image_part_when_both_branches_add_the_same_picture(self):
        base = empty_base()
        left, right = branches(base)

        left.add_picture(image_path("python-powered.png"))
        right.add_picture(image_path("python-powered.png"))

        result = merge_documents(base, fork(left), fork(right))
        doc = fork(result.document)

        rids = doc._element.body.xpath(".//a:blip/@r:embed")
        # -- identical insertion is taken once, and the blob is deduplicated --
        assert len(rids) == 1
        assert doc.part.rels[rids[0]].target_part.blob == image_blob(
            "python-powered.png"
        )

    def it_preserves_a_picture_inside_a_table_cell_after_save_and_reopen(self):
        base_doc = Document()
        table = base_doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "x"
        table.cell(0, 1).text = "y"
        base = fork(base_doc)
        left, right = branches(base)

        run = left.tables[0].cell(0, 0).paragraphs[0].add_run()
        run.add_picture(image_path("python-icon.jpeg"))
        right.tables[0].cell(0, 1).text = "Y"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        doc = fork(result.document)
        embeds = doc.tables[0]._tbl.xpath(".//a:blip/@r:embed")
        assert len(embeds) == 1
        assert doc.part.rels[embeds[0]].target_part.blob == image_blob(
            "python-icon.jpeg"
        )
        assert doc.tables[0].cell(0, 1).text == "Y"


class DescribeHyperlinkReferenceMerge:
    def it_remaps_hyperlinks_sharing_a_source_rId_to_their_own_urls(self):
        base = empty_base()
        left, right = branches(base)

        add_hyperlink_paragraph(left, "https://example.com/left", "Left Link")
        add_hyperlink_paragraph(right, "https://example.com/right", "Right Link")
        assert (
            left._element.body.xpath(".//w:hyperlink/@r:id")[0]
            == right._element.body.xpath(".//w:hyperlink/@r:id")[0]
        )

        result = merge_documents(base, fork(left), fork(right))
        doc = fork(result.document)

        links = [
            (
                "".join(t.text or "" for t in link.iter(qn("w:t"))),
                doc.part.rels[link.get(qn("r:id"))].target_ref,
            )
            for link in doc._element.body.xpath(".//w:hyperlink")
        ]
        assert ("Left Link", "https://example.com/left") in links
        assert ("Right Link", "https://example.com/right") in links
        rids = [link.get(qn("r:id")) for link in doc._element.body.xpath(".//w:hyperlink")]
        assert len(set(rids)) == 2


class DescribeNumberingReferenceMerge:
    def it_remaps_numbering_definitions_that_share_a_source_numId(self):
        base = empty_base()
        left, right = branches(base)

        add_numbered_paragraph(left, "l-item", 42, 42, "decimal")
        add_numbered_paragraph(right, "r-item", 42, 43, "lowerRoman")

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        doc = fork(result.document)
        num_ids = doc._element.body.xpath(".//w:numPr/w:numId/@w:val")
        assert len(set(num_ids)) == 2
        formats = {numbering_format_for(doc, num_id) for num_id in num_ids}
        assert formats == {"decimal", "lowerRoman"}
        assert paragraph_texts(doc) == ["l-item", "r-item"]


class DescribeStyleReferenceMerge:
    def it_duplicates_a_style_id_whose_definition_differs_and_rewires_references(self):
        base = empty_base()
        left, right = branches(base)

        left_style = left.styles.add_style("MyStyle", WD_STYLE_TYPE.PARAGRAPH)
        left_style.font.size = Pt(20)
        right_style = right.styles.add_style("MyStyle", WD_STYLE_TYPE.PARAGRAPH)
        right_style.font.size = Pt(40)
        left.add_paragraph("left-styled").style = left_style
        right.add_paragraph("right-styled").style = right_style

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        doc = fork(result.document)
        by_id = {
            style.style_id: style.font.size
            for style in doc.styles
            if style.style_id == "MyStyle"
            or style.style_id.startswith("MyStyleMerge")
        }
        assert "MyStyle" in by_id
        assert any(style_id.startswith("MyStyleMerge") for style_id in by_id)
        used_ids = {
            paragraph.text: paragraph.style.style_id
            for paragraph in doc.paragraphs
            if paragraph.text in ("left-styled", "right-styled")
        }
        assert by_id[used_ids["left-styled"]] == Pt(20)
        assert by_id[used_ids["right-styled"]] == Pt(40)


class DescribeMergeDeterminism:
    def it_produces_byte_identical_packages_and_conflicts_when_run_twice(self):
        base = empty_base()
        left, right = branches(base)

        left.add_picture(image_path("python-powered.png"))
        add_hyperlink_paragraph(right, "https://example.com/x", "X")

        inputs = (base, fork(left), fork(right))
        first = merge_documents(*(fork(doc) for doc in inputs))
        second = merge_documents(*(fork(doc) for doc in inputs))

        assert first.conflicts == second.conflicts
        first_parts = zip_part_bytes(first.document)
        second_parts = zip_part_bytes(second.document)
        assert set(first_parts) == set(second_parts)
        assert first_parts == second_parts

    def it_can_remerge_its_own_result_without_losing_or_duplicating_content(self):
        base = fork(Document())
        base.add_paragraph("one")
        base = fork(base)
        left, right = branches(base)

        left.add_picture(image_path("python-powered.png"))
        add_hyperlink_paragraph(right, "https://example.com/x", "X")

        merged = merge_documents(base, fork(left), fork(right))
        assert not merged.has_conflicts
        merged_doc = fork(merged.document)

        # -- feeding the merged document back in as both branches is a clean merge --
        again = merge_documents(base, fork(merged_doc), fork(merged_doc))

        assert not again.has_conflicts, again.conflicts
        doc = fork(again.document)
        assert paragraph_texts(doc) == paragraph_texts(merged_doc)
        assert len(doc._element.body.xpath(".//a:blip")) == 1
        links = doc._element.body.xpath(".//w:hyperlink")
        assert len(links) == 1
        assert doc.part.rels[links[0].get(qn("r:id"))].target_ref == (
            "https://example.com/x"
        )
