"""Shared builders for the three-way document merge tests.

Every scenario follows the same shape: build a *base* document, fork it into two
independent branches, edit the branches, fork them once more (to simulate a save and
reopen between editing and merging), then merge.
"""

from __future__ import annotations

import io
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import parse_xml

from ..unitutil.file import test_file


def fork(doc: Document) -> Document:
    """Round-trip `doc` through the package format, like closing and reopening it."""
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return Document(stream)


def text_document(*texts: str) -> Document:
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    return fork(doc)


def branches(doc: Document) -> tuple[Document, Document]:
    return fork(doc), fork(doc)


def delete_paragraph(doc: Document, index: int) -> None:
    paragraph = doc.paragraphs[index]._p
    paragraph.getparent().remove(paragraph)


def add_hyperlink_paragraph(doc: Document, url: str, text: str) -> None:
    r_id = doc.part.rels.get_or_add_ext_rel(RT.HYPERLINK, url)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    doc.add_paragraph()._p.append(hyperlink)


def add_numbered_paragraph(
    doc: Document, text: str, num_id: int, abstract_num_id: int, num_fmt: str
) -> None:
    """Add a paragraph that references a freshly-built numbering definition."""
    root = doc.part.numbering_part.element
    abstract_xml = (
        '<w:abstractNum %s w:abstractNumId="%d">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/>'
        '<w:numFmt w:val="%s"/><w:lvlText w:val="%%1."/></w:lvl>'
        "</w:abstractNum>"
    ) % (nsdecls("w"), abstract_num_id, num_fmt)
    num_xml = (
        '<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>'
        % (nsdecls("w"), num_id, abstract_num_id)
    )
    root.append(parse_xml(abstract_xml))
    root.append(parse_xml(num_xml))

    paragraph = doc.add_paragraph(text)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    paragraph._p.get_or_add_pPr().append(num_pr)


def apply_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def numbering_format_for(doc: Document, num_id: str) -> str:
    root = doc.part.numbering_part.element
    abstract_ids = root.xpath(
        './w:num[@w:numId="%s"]/w:abstractNumId/@w:val' % num_id
    )
    formats = root.xpath(
        './w:abstractNum[@w:abstractNumId="%s"]//w:numFmt/@w:val' % abstract_ids[0]
    )
    return formats[0]


def image_path(name: str) -> str:
    return test_file(name)


def paragraph_texts(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


def cell_texts(doc: Document, table_index: int = 0) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in doc.tables[table_index].rows]


def zip_part_bytes(doc: Document) -> dict[str, bytes]:
    """Map every package part name to its bytes (for determinism comparisons)."""
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        return {name: archive.read(name) for name in archive.namelist()}
