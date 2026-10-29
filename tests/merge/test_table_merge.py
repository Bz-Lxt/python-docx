# pyright: reportPrivateUsage=false

"""Integration tests for three-way merging of tables, including merged cells."""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.merge import merge_documents
from docx.shared import Inches

from .helpers import (
    apply_shading,
    branches,
    cell_texts,
    fork,
    paragraph_texts,
)


def table_document(rows: int, cols: int, prefix: str = "b") -> Document:
    doc = Document()
    table = doc.add_table(rows=rows, cols=cols)
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = "%s%d%d" % (prefix, r, c)
    return fork(doc)


class DescribeTableMerge:
    def it_combines_a_row_insertion_with_an_unrelated_cell_edit(self):
        base = table_document(2, 2)
        left, right = branches(base)

        left.tables[0].add_row().cells[0].text = "new-row"
        right.tables[0].cell(1, 1).text = "r11"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        assert cell_texts(result.document) == [
            ["b00", "b01"],
            ["b10", "r11"],
            ["new-row", ""],
        ]

    def it_combines_a_row_deletion_with_an_unrelated_cell_edit(self):
        base = table_document(3, 1)
        left, right = branches(base)

        tr = left.tables[0]._tbl.tr_lst[1]
        tr.getparent().remove(tr)
        right.tables[0].cell(0, 0).text = "top-right"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        assert cell_texts(result.document) == [["top-right"], ["b20"]]

    def it_locates_a_cell_content_conflict_at_the_cell_path(self):
        base = table_document(2, 2)
        left, right = branches(base)

        left.tables[0].cell(0, 0).text = "LEFT"
        right.tables[0].cell(0, 0).text = "RIGHT"

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert conflict.location == "body[0]/row[0]/cell[0][0]"
        assert conflict.kind == "paragraph"
        assert conflict.branches == ("left", "right")
        assert conflict.base_text == "b00"
        assert conflict.left_text == "LEFT"
        assert conflict.right_text == "RIGHT"
        merged_cell = result.document.tables[0].cell(0, 0)
        assert [p.text for p in merged_cell.paragraphs] == [
            "<<<<<<< left",
            "LEFT",
            "=======",
            "RIGHT",
            ">>>>>>> right",
        ]
        # -- the rest of the table is untouched --
        table = result.document.tables[0]
        assert table.cell(0, 1).text == "b01"
        assert table.cell(1, 0).text == "b10"
        assert table.cell(1, 1).text == "b11"

    def it_merges_cell_shading_on_one_side_with_text_edit_on_the_other(self):
        base = table_document(1, 1)
        left, right = branches(base)

        apply_shading(left.tables[0].cell(0, 0), "FF0000")
        right.tables[0].cell(0, 0).text = "both-kept"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        doc = fork(result.document)
        cell = doc.tables[0].cell(0, 0)
        assert cell.text == "both-kept"
        assert cell._tc.xpath(".//w:shd/@w:fill") == ["FF0000"]

    def it_reports_conflicting_cell_shading_as_a_cell_properties_conflict(self):
        base = table_document(1, 1)
        left, right = branches(base)

        apply_shading(left.tables[0].cell(0, 0), "FF0000")
        apply_shading(right.tables[0].cell(0, 0), "00FF00")

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert conflict.location == "body[0]/row[0]/cell[0]/properties"
        assert conflict.kind == "table-cell-properties"
        assert conflict.branches == ("left", "right")

    def it_merges_vertically_merged_cells_property_and_content_independently(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 0).text = "merged"
        table.cell(0, 1).text = "topright"
        table.cell(1, 1).text = "botright"
        base = fork(doc)
        left, right = branches(base)

        apply_shading(left.tables[0].cell(0, 0), "0000FF")
        right.tables[0].cell(0, 0).text = "MERGED-RIGHT"
        right.tables[0].cell(1, 1).text = "BOTTOM-RIGHT"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts, result.conflicts
        doc = fork(result.document)
        anchor = doc.tables[0].cell(0, 0)
        assert anchor.text == "MERGED-RIGHT"
        assert anchor._tc.xpath(".//w:shd/@w:fill") == ["0000FF"]
        assert doc.tables[0].cell(1, 1).text == "BOTTOM-RIGHT"

    def it_merges_horizontally_merged_cells_property_and_content(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "wide"
        base = fork(doc)
        left, right = branches(base)

        apply_shading(left.tables[0].cell(0, 0), "AAAAAA")
        right.tables[0].cell(0, 0).text = "WIDE-RIGHT"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts, result.conflicts
        doc = fork(result.document)
        cell = doc.tables[0].cell(0, 0)
        assert cell.text == "WIDE-RIGHT"
        assert cell._tc.xpath(".//w:shd/@w:fill") == ["AAAAAA"]

    def it_reports_a_single_block_conflict_when_column_counts_diverge(self):
        base = table_document(1, 2)
        left, right = branches(base)

        left.tables[0].add_column(Inches(1))
        right.tables[0].cell(0, 0).text = "edited"

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        assert result.conflicts[0].kind == "table"
        assert result.conflicts[0].location == "body[0]"
        # -- the summary carries the whole right-side table's text --
        assert result.conflicts[0].right_text == "edited b01"

    def it_reports_row_delete_against_row_edit_as_a_table_row_conflict(self):
        base = table_document(2, 1)
        left, right = branches(base)

        tr = left.tables[0]._tbl.tr_lst[0]
        tr.getparent().remove(tr)
        right.tables[0].cell(0, 0).text = "R0-right"

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert conflict.kind == "table-row"
        assert conflict.location == "body[0]/row[0]"
        assert conflict.left_text == ""
        assert conflict.right_text == "R0-right"
        # -- the deleted-vs-edited row renders a complete, ordered marker block --
        cell = result.document.tables[0].cell(0, 0)
        assert [p.text for p in cell.paragraphs] == [
            "<<<<<<< left",
            "=======",
            "R0-right",
            ">>>>>>> right",
        ]
        # -- the unaffected second row survives --
        assert result.document.tables[0].cell(1, 0).text == "b10"

    def it_renders_a_complete_marker_block_when_right_deletes_an_edited_row(self):
        base = table_document(2, 1)
        left, right = branches(base)

        left.tables[0].cell(0, 0).text = "L0-left"
        tr = right.tables[0]._tbl.tr_lst[0]
        tr.getparent().remove(tr)

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert conflict.kind == "table-row"
        assert conflict.location == "body[0]/row[0]"
        assert conflict.left_text == "L0-left"
        assert conflict.right_text == ""
        cell = result.document.tables[0].cell(0, 0)
        assert [p.text for p in cell.paragraphs] == [
            "<<<<<<< left",
            "L0-left",
            "=======",
            ">>>>>>> right",
        ]
        assert result.document.tables[0].cell(1, 0).text == "b10"

    def it_merges_nested_tables_cell_by_cell(self):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "i0"
        inner.cell(0, 1).text = "i1"
        base = fork(doc)
        left, right = branches(base)

        left.tables[0].cell(0, 0).tables[0].cell(0, 0).text = "LEFT-INNER"
        right.tables[0].cell(0, 0).tables[0].cell(0, 1).text = "RIGHT-INNER"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        inner_result = result.document.tables[0].cell(0, 0).tables[0]
        assert inner_result.cell(0, 0).text == "LEFT-INNER"
        assert inner_result.cell(0, 1).text == "RIGHT-INNER"

    def it_takes_a_table_property_change_made_on_only_one_branch(self):
        base = table_document(1, 1)
        left, right = branches(base)

        left.tables[0].alignment = WD_TABLE_ALIGNMENT.CENTER

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        assert (
            result.document.tables[0].alignment == WD_TABLE_ALIGNMENT.CENTER
        )

    def it_reports_conflicting_table_properties_without_losing_content(self):
        base = table_document(1, 1)
        left, right = branches(base)

        left.tables[0].alignment = WD_TABLE_ALIGNMENT.CENTER
        right.tables[0].alignment = WD_TABLE_ALIGNMENT.RIGHT

        result = merge_documents(base, fork(left), fork(right))

        assert len(result.conflicts) == 1
        conflict = result.conflicts[0]
        assert conflict.kind == "table-properties"
        assert conflict.location == "body[0]"
        assert result.document.tables[0].cell(0, 0).text == "b00"

    def it_preserves_cells_when_one_side_only_changes_grid_widths(self):
        base = table_document(1, 2)
        left, right = branches(base)

        left.tables[0].columns[0].width = Inches(3)
        right.tables[0].cell(0, 1).text = "edited"

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts, result.conflicts
        assert result.document.tables[0].cell(0, 1).text == "edited"

    def it_keeps_unchanged_tables_completely_untouched(self):
        base = table_document(2, 2)
        left, right = branches(base)
        left.add_paragraph("after-table")

        result = merge_documents(base, fork(left), fork(right))

        assert not result.has_conflicts
        assert cell_texts(result.document) == [
            ["b00", "b01"],
            ["b10", "b11"],
        ]
        assert paragraph_texts(result.document) == ["after-table"]
        # -- no stray relationship references introduced for a plain table --
        assert not result.document._element.body.xpath(".//a:blip")
        assert not result.document._element.body.xpath(".//w:hyperlink")
